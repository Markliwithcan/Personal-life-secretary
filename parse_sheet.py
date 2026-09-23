"""Read untrusted plan files. No formula execution, macros, or external links."""
import csv
from datetime import datetime, date, time
import io
import json
from pathlib import Path
import sys
import zipfile
from openpyxl import load_workbook
from openpyxl.styles.numbers import is_datetime

ALIASES={
 'title': ['事项','任务','任务名称','待办','待办事项','计划','内容','title','task'],
 'due': ['日期','截止日期','截止时间','到期时间','计划日期','due','date'],
 'remind_at': ['提醒时间','提醒日期','remind_at','reminder'],
 'priority': ['优先级','重要程度','priority'],
 'notes': ['备注','说明','notes'],
 'status': ['状态','完成状态','status'],
 'recurrence': ['重复','周期','循环','recurrence'],
}
LOOKUP={alias:field for field,aliases in ALIASES.items() for alias in aliases}

def scalar(value):
    if isinstance(value,(datetime,date,time)): return value.isoformat()
    return str(value or '').strip()

def date_value(value):
    if value is None or value=='':return None
    if isinstance(value,datetime):return value.isoformat()
    if isinstance(value,date):return value.isoformat()
    text=str(value).strip().replace('/','-').replace('年','-').replace('月','-').replace('日','')
    for fmt in ('%Y-%m-%d','%Y-%m-%d %H:%M','%Y-%m-%d %H:%M:%S'):
        try:
            parsed=datetime.strptime(text,fmt)
            return parsed.date().isoformat() if fmt=='%Y-%m-%d' else parsed.isoformat()
        except ValueError:pass
    try:return datetime.fromisoformat(text).isoformat()
    except ValueError:raise ValueError('日期需写明年月日；不能只写时刻或未经格式化的数字')

def parse(filename,default_reminder_time='09:00'):
    path=Path(filename)
    if path.stat().st_size>10*1024*1024:raise ValueError('文件超过 10MB')
    raw=path.read_bytes()
    workbook=None
    if raw[:2]==b'PK':
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(x.file_size for x in z.infolist())>50*1024*1024 or len(z.infolist())>2000:raise ValueError('工作簿解压后过大')
            if any('vbaproject' in n.lower() for n in z.namelist()):raise ValueError('不支持带宏的工作簿')
        workbook=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        if len(workbook.worksheets)>10:
            workbook.close();raise ValueError('最多支持 10 张工作表，请拆分文件')
        def values(ws):
            for row in ws.iter_rows():
                yield [cell.value.date() if isinstance(cell.value,datetime) and is_datetime(cell.number_format)=='date' else cell.value for cell in row]
        sheets=[(ws.title,values(ws)) for ws in workbook.worksheets]
    elif raw.startswith(b'\xd0\xcf'):
        raise ValueError('暂不支持旧版 .xls，请另存为 .xlsx')
    else:
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:text=raw.decode('gb18030')
        if '\x00' in text:raise ValueError('不是可识别的计划表')
        try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=',\t;')
        except csv.Error:dialect=csv.excel
        sheets=[('CSV',csv.reader(io.StringIO(text),dialect))]
    tasks,warnings=[],[]
    try:
        for name,rows in sheets:
            headers=None
            for index,row in enumerate(rows,1):
                if index>2001:raise ValueError('每张表最多支持 2000 行')
                row=list(row)
                if len(row)>100:raise ValueError('每张表最多支持 100 列')
                if headers is None:
                    candidate={i:LOOKUP.get(scalar(v).lower()) for i,v in enumerate(row)}
                    if 'title' in candidate.values():headers=candidate;continue
                    if index>=10:break
                    continue
                data={field:row[i] for i,field in headers.items() if field and i<len(row)}
                title=scalar(data.get('title'))
                if not title:continue
                if scalar(data.get('status')).lower() in ['完成','已完成','取消','已取消','done','completed','cancelled']:continue
                if any(isinstance(v,str) and v.startswith('=') for v in data.values()):
                    warnings.append(f'{name} 第{index}行含公式，未导入；请将公式结果粘贴为值。');continue
                if len(title)>200:warnings.append(f'{name} 第{index}行标题过长，未导入。');continue
                try:
                    due=date_value(data.get('due'));remind=date_value(data.get('remind_at'))
                    if remind and len(remind)==10:remind+='T'+default_reminder_time+':00';warnings.append(f'{name} 第{index}行提醒只有日期，预设当天 {default_reminder_time}，请在导入预览核对。')
                    priority=scalar(data.get('priority')).lower()
                    recurrence=scalar(data.get('recurrence')).lower()
                    if recurrence and recurrence not in ['每天','每日','daily','每周','weekly','无','不重复','none']:raise ValueError('仅支持每天/每周循环')
                    tasks.append({'title':title,'due':due,'remind_at':remind,'priority':'high' if priority in ['高','重要','紧急','high','p1'] else 'low' if priority in ['低','low','p3'] else 'normal',
                                  'recurrence':'daily' if recurrence in ['每天','每日','daily'] else 'weekly' if recurrence in ['每周','weekly'] else None,
                                  'notes':(scalar(data.get('notes'))+f'\n来源：{name} 第{index}行')[:2000]})
                    if len(tasks)>300:raise ValueError('一次最多导入 300 个待办，请拆分文件')
                except ValueError as e:warnings.append(f'{name} 第{index}行未导入：{e}')
            if headers is None:warnings.append(f'{name} 未找到“事项/任务/title”等标题列。')
    finally:
        if workbook:workbook.close()
    if len(tasks)>300:raise ValueError('一次最多导入 300 个待办，请拆分文件')
    return {'tasks':tasks,'warnings':warnings}

if __name__=='__main__':
    try: print(json.dumps(parse(sys.argv[1],sys.argv[2] if len(sys.argv)>2 else '09:00'),ensure_ascii=False))
    except Exception as e:
        print(json.dumps({'error':str(e)[:300]},ensure_ascii=False));sys.exit(1)
