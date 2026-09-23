import csv
from datetime import datetime
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook
from parse_sheet import parse

class ParseTests(unittest.TestCase):
    def test_csv_dates_status_and_formula(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'test.csv'
            with p.open('w',encoding='utf-8-sig',newline='') as f:
                w=csv.writer(f);w.writerows([['事项','日期','状态'],['买菜','2026/9/17','未完成'],['已做','2026/9/17','已完成'],['=HYPERLINK("http://bad")','2026/9/17',''],['坏日期','9月17日','']])
            result=parse(p)
            self.assertEqual(len(result['tasks']),1)
            self.assertEqual(result['tasks'][0]['due'],'2026-09-17')
            self.assertEqual(len(result['warnings']),2)

    def test_xlsx_native_dates_and_source(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'test.xlsx';wb=Workbook();ws=wb.active;ws.title='九月计划'
            ws.append(['任务','截止时间','提醒时间','优先级']);ws.append(['报告',datetime(2026,9,18,18),datetime(2026,9,18,9),'高']);wb.save(p);wb.close()
            result=parse(p);self.assertEqual(result['tasks'][0]['priority'],'high');self.assertIn('九月计划 第2行',result['tasks'][0]['notes'])

    def test_legacy_xls_rejected(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'test.xls';p.write_bytes(b'\xd0\xcf\x11\xe0')
            with self.assertRaisesRegex(ValueError,'另存为'):parse(p)

    def test_date_only_excel_uses_configured_default_time(self):
        with tempfile.TemporaryDirectory() as d:
            p=Path(d)/'test.xlsx';wb=Workbook();ws=wb.active
            ws.append(['事项','日期','提醒时间']);ws.append(['报告',datetime(2026,9,18),datetime(2026,9,17)])
            ws['B2'].number_format='yyyy-mm-dd';ws['C2'].number_format='yyyy-mm-dd';wb.save(p);wb.close()
            result=parse(p,'10:30');self.assertEqual(result['tasks'][0]['due'],'2026-09-18');self.assertEqual(result['tasks'][0]['remind_at'],'2026-09-17T10:30:00')

if __name__=='__main__':unittest.main()
